"""Read the master resume (.docx / .pdf / .txt / .md) into a ``MasterProfile``.

Parsed once and cached to JSON. That cache is the fact base every later stage is
constrained against, so it's worth spending a good model on: a skill the parser
misses is a skill the tailored resume can never legitimately claim.
"""

from __future__ import annotations

import logging
from pathlib import Path

from ..config import Config
from ..llm import Claude, LLMError
from ..models import MasterProfile
from ..util import clean_text

log = logging.getLogger(__name__)

PARSE_INSTRUCTIONS = """\
Convert this resume into structured data.

Rules:

- Transcribe, don't improve. Keep bullets close to their original wording;
  fixing obvious typos is fine, rewriting is not.
- skills: collect every technology, tool, language, framework, library, platform,
  and technical competency mentioned anywhere in the document — including ones
  that only appear inside a project description, a certification title, or a
  course name. This list is what later stages are allowed to draw on, so a skill
  omitted here is effectively erased.
- headline: the role line under the name, or the target role the resume is
  aimed at.
- experience: internships and part-time roles count. Keep dates exactly as
  written.
- projects: personal and academic projects, with their tech stacks.
- education: one line per qualification, preformatted with " | " between the
  parts, e.g. "Institution Name | B.Tech CSE (AI & ML) | CGPA 7.04/10 | 2023–2027".
- activities: leadership positions, clubs, volunteering.
- Leave a field empty rather than guessing at it.
"""


def extract_text(path: Path) -> str:
    """Pull plain text out of a resume file."""
    if not path.exists():
        raise FileNotFoundError(
            f"Master resume not found at {path}. Point `profile.master_resume` "
            "in config.yaml at your resume, or run: job-agent init <path>"
        )

    suffix = path.suffix.lower()
    if suffix == ".docx":
        return clean_text(_from_docx(path))
    if suffix == ".pdf":
        return clean_text(_from_pdf(path))
    if suffix in (".txt", ".md"):
        return clean_text(path.read_text(encoding="utf-8", errors="replace"))
    raise ValueError(f"Unsupported resume format: {suffix} (use .docx, .pdf, .txt, .md)")


def _from_docx(path: Path) -> str:
    from docx import Document

    doc = Document(str(path))
    parts = [p.text for p in doc.paragraphs]
    # Many resume templates lay out contact rows and dates in tables.
    for table in doc.tables:
        for row in table.rows:
            cells = [c.text.strip() for c in row.cells if c.text.strip()]
            if cells:
                parts.append(" | ".join(cells))
    return "\n".join(parts)


def _from_pdf(path: Path) -> str:
    from pypdf import PdfReader

    reader = PdfReader(str(path))
    return "\n".join(page.extract_text() or "" for page in reader.pages)


def load_profile(
    claude: Claude | None, config: Config, force_reparse: bool = False
) -> MasterProfile:
    """Return the structured master profile, parsing and caching on first use.

    ``claude`` may be None when a valid cache exists — that lets ``--dry-run``
    work without an API key.
    """
    cache = config.profile_cache_path

    if cache.exists() and not force_reparse:
        source = config.master_resume_path
        if not source.exists() or cache.stat().st_mtime >= source.stat().st_mtime:
            try:
                return MasterProfile.model_validate_json(
                    cache.read_text(encoding="utf-8")
                )
            except Exception as exc:  # corrupt cache shouldn't be fatal
                log.warning("profile cache unreadable (%s) — re-parsing", exc)

    if claude is None:
        raise LLMError(
            "The master resume needs parsing, which requires an Anthropic API "
            "key. Add ANTHROPIC_API_KEY to .env and run: job-agent init <resume>"
        )

    text = extract_text(config.master_resume_path)
    log.info("parsing master resume (%d characters)", len(text))

    profile = claude.parse(
        model=config.models.resume_parse,
        schema=MasterProfile,
        instructions=PARSE_INSTRUCTIONS,
        user_content=text,
        max_tokens=12000,
        effort=config.models.effort.resume_parse,
    )

    cache.parent.mkdir(parents=True, exist_ok=True)
    cache.write_text(profile.model_dump_json(indent=2), encoding="utf-8")
    log.info(
        "parsed: %d skills, %d roles, %d projects",
        len(profile.skills), len(profile.experience), len(profile.projects),
    )
    return profile
